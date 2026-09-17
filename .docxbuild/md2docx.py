#!/usr/bin/env python
"""ARCHITECTURE.md -> a styled Word .docx, fully offline.

Phase 1 (--extract): writes every ```mermaid block to diagrams/dNN.mmd
Phase 2 (default):   builds the .docx, embedding diagrams/dNN.png when present
                     and falling back to the Mermaid source as a code block.
Both phases number diagrams by order of appearance, so they stay in step.
"""
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

ROOT = r"D:\OptiDesk"
SRC = os.path.join(ROOT, "ARCHITECTURE.md")
OUT = os.path.join(ROOT, "OptiDesk-Architecture.docx")
DIAG = os.path.join(ROOT, ".docxbuild", "diagrams")

ACCENT = RGBColor(0x1F, 0x38, 0x64)
ACCENT2 = RGBColor(0x2F, 0x54, 0x96)
CODEFG = RGBColor(0xA3, 0x20, 0x20)
GREY = RGBColor(0x5A, 0x5A, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HDR_FILL = "1F3864"
ZEBRA = "F2F6FB"
CODE_FILL = "F7F7F8"
INL_FILL = "F0F0F1"
CALL_FILL = "FFF7E6"
MONO = "Consolas"
MAX_IMG = Cm(17.2)
MAX_IMG_H = Cm(22.5)


def shade(el, fill):
    """w:shd — valid on tcPr (before tcMar) and at the tail of rPr."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el.append(shd)


def preserve(par):
    """Leading/trailing spaces in code lines survive only with xml:space."""
    for t in par._p.iter(qn("w:t")):
        t.set(qn("xml:space"), "preserve")


def set_mono(run):
    fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), MONO)


def tbl_borders(table, color="D9D9D9", sz=6):
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)
    pr = table._tbl.tblPr
    for old in pr.findall(qn("w:tblBorders")):
        pr.remove(old)
    # Schema order: tblBorders comes before shd/tblLayout/tblCellMar/tblLook.
    pr.insert_element_before(b, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")


def cell_margins(cell, twips):
    mar = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + side)
        e.set(qn("w:w"), str(twips))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    cell._tc.get_or_add_tcPr().append(mar)


def repeat_header(row):
    """Header row reprints at the top of every page the table spills onto."""
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(h)


INLINE = re.compile(r"(`[^`]+`|\*\*.+?\*\*|(?<![\*\w])\*[^*\s][^*]*\*(?!\*))")


def add_inline(par, text, size=10.0, bold=False, mono_ok=True):
    """Render `code`, **bold** and *italic* spans as separate runs."""
    for part in INLINE.split(text):
        if not part:
            continue
        if len(part) > 1 and part[0] == "`" and part[-1] == "`":
            r = par.add_run(part[1:-1])
            set_mono(r)
            r.font.size = Pt(size - 0.5)
            r.bold = bold
            if mono_ok:
                r.font.color.rgb = CODEFG
                shade(r._r.get_or_add_rPr(), INL_FILL)
        elif part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            r.bold = True
            r.font.size = Pt(size)
        elif len(part) > 1 and part[0] == "*" and part[-1] == "*":
            r = par.add_run(part[1:-1])
            r.italic = True
            r.bold = bold
            r.font.size = Pt(size)
        else:
            r = par.add_run(part)
            r.bold = bold
            r.font.size = Pt(size)
    preserve(par)
    return par


def split_row(line):
    """Split a table row on | — but never on an escaped \\| (GFM's way of
    putting a literal pipe in a cell) nor on a | inside a `code span`."""
    line = line.strip().replace("\\|", "\x00")
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    cells, buf, in_code = [], [], False
    for ch in line:
        if ch == "`":
            in_code = not in_code
            buf.append(ch)
        elif ch == "|" and not in_code:
            cells.append("".join(buf))
            buf = []
        else:
            buf.append(ch)
    cells.append("".join(buf))
    return [c.replace("\x00", "|").strip() for c in cells]


SEP = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?$")
HEAD = re.compile(r"^(#{1,4})\s+(.*)$")
NUM = re.compile(r"^(\d+)\.\s+(.*)$")
STOP = ("|", "```", "#", "- ", "> ", "---")


def parse(md):
    lines = md.split("\n")
    blocks, i, n = [], 0, len(lines)
    while i < n:
        s = lines[i].strip()
        if s.startswith("```"):
            lang = s[3:].strip().lower()
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            kind = "mermaid" if lang == "mermaid" else "code"
            blocks.append((kind, "\n".join(buf)))
            continue
        if s.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                if not SEP.match(lines[i].strip()):
                    rows.append(split_row(lines[i]))
                i += 1
            blocks.append(("table", rows))
            continue
        m = HEAD.match(s)
        if m:
            blocks.append(("h%d" % len(m.group(1)), m.group(2).strip()))
            i += 1
            continue
        if s.startswith("> "):
            buf = []
            while i < n and lines[i].strip().startswith("> "):
                buf.append(lines[i].strip()[2:])
                i += 1
            blocks.append(("quote", " ".join(buf)))
            continue
        if s.startswith("- "):
            buf = []
            while i < n and lines[i].strip().startswith("- "):
                buf.append(lines[i].strip()[2:])
                i += 1
            blocks.append(("bullets", buf))
            continue
        if NUM.match(s):
            buf = []
            while i < n and NUM.match(lines[i].strip()):
                mm = NUM.match(lines[i].strip())
                buf.append((mm.group(1), mm.group(2)))
                i += 1
            blocks.append(("numbers", buf))
            continue
        if not s or s.startswith("---"):
            i += 1
            continue
        buf = []
        while i < n:
            t = lines[i].strip()
            if not t or t.startswith(STOP) or NUM.match(t):
                break
            buf.append(t)
            i += 1
        blocks.append(("para", " ".join(buf)))
    return blocks


def extract(blocks):
    os.makedirs(DIAG, exist_ok=True)
    k = 0
    for kind, payload in blocks:
        if kind != "mermaid":
            continue
        k += 1
        path = os.path.join(DIAG, "d%02d.mmd" % k)
        with open(path, "w", encoding="utf-8") as fh:
            fh.write(payload.strip() + "\n")
    print("extracted %d mermaid diagrams -> %s" % (k, DIAG))


def spacer(doc, pts=5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(pts)
    pf.line_spacing = 1.0
    p.add_run("").font.size = Pt(2)


def code_block(doc, code, size=8.0, fill=CODE_FILL, border="E3E3E6"):
    """A 1x1 shaded table — the only reliable boxed code block in OOXML."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    tbl_borders(t, border, 4)
    cell = t.cell(0, 0)
    shade(cell._tc.get_or_add_tcPr(), fill)
    cell_margins(cell, 140)
    p = cell.paragraphs[0]
    pf = p.paragraph_format
    pf.space_before = Pt(3)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.0
    r = p.add_run()
    set_mono(r)
    r.font.size = Pt(size)
    for idx, line in enumerate(code.rstrip("\n").split("\n")):
        if idx:
            r.add_break()
        r.add_text(line)
    preserve(p)
    spacer(doc)


def callout(doc, text):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    tbl_borders(t, "E8C77A", 4)
    cell = t.cell(0, 0)
    shade(cell._tc.get_or_add_tcPr(), CALL_FILL)
    cell_margins(cell, 140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline(p, text, size=9.5)
    spacer(doc)


def diagram(doc, index, caption):
    png = os.path.join(DIAG, "d%02d.png" % index)
    if not os.path.exists(png):
        return False
    img = DocxImage.from_file(png)
    # Cap on both axes: the erDiagram and the sequenceDiagram are taller than
    # one page at full column width, and Word will not shrink them for us.
    width = min(int(img.width), int(MAX_IMG))
    aspect = img.px_height / float(img.px_width)
    if width * aspect > int(MAX_IMG_H):
        width = int(int(MAX_IMG_H) / aspect)
    doc.add_picture(png, width=Emu(width))
    pic = doc.paragraphs[-1]
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.space_before = Pt(4)
    pic.paragraph_format.space_after = Pt(2)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(9)
    r = cap.add_run("Diagram %d — %s" % (index, caption))
    r.italic = True
    r.font.size = Pt(8)
    r.font.color.rgb = GREY
    return True


def table_block(doc, rows):
    if not rows:
        return
    hdr = rows[0]
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=1, cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = True
    tbl_borders(t)
    for i in range(cols):
        c = t.rows[0].cells[i]
        shade(c._tc.get_or_add_tcPr(), HDR_FILL)
        cell_margins(c, 90)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        add_inline(p, hdr[i] if i < len(hdr) else "", 8.5, True, mono_ok=False)
        for r in p.runs:
            r.bold = True
            r.font.color.rgb = WHITE
    repeat_header(t.rows[0])
    for ri, row in enumerate(rows[1:]):
        cells = t.add_row().cells
        for i in range(cols):
            c = cells[i]
            if ri % 2 == 0:
                shade(c._tc.get_or_add_tcPr(), ZEBRA)
            cell_margins(c, 90)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_inline(p, row[i] if i < len(row) else "", 8.5)
    spacer(doc)


def heading(doc, level, text):
    if level == 1:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(23)
        r.font.color.rgb = ACCENT
        return
    p = doc.add_paragraph(style="Heading %d" % (level - 1))
    pf = p.paragraph_format
    pf.page_break_before = level == 2
    pf.space_before = Pt(10 if level == 2 else 12)
    pf.space_after = Pt(5)
    pf.keep_with_next = True
    add_inline(p, text, 15.0 if level == 2 else 11.5, mono_ok=False)
    for r in p.runs:
        r.bold = True
        r.font.color.rgb = ACCENT if level == 2 else ACCENT2


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.08
        add_inline(p, it)


def numbers(doc, items):
    """Literal numbers + hanging indent: Word's List Number style would keep
    counting across the whole document, and this file has nine separate lists."""
    for num, txt in items:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.85)
        pf.first_line_indent = Cm(-0.85)
        pf.space_after = Pt(2)
        pf.line_spacing = 1.08
        r = p.add_run("%s.\t" % num)
        r.bold = True
        add_inline(p, txt)


def subtitle(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(
        "MERN  ·  Clerk  ·  Gemini RAG  ·  Socket.io      "
        "generated from ARCHITECTURE.md"
    )
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = GREY


def toc(doc, blocks):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("Contents")
    r.bold = True
    r.font.size = Pt(12.5)
    r.font.color.rgb = ACCENT
    for kind, payload in blocks:
        if kind not in ("h2", "h3"):
            continue
        q = doc.add_paragraph()
        pf = q.paragraph_format
        pf.left_indent = Cm(0.3 if kind == "h2" else 1.1)
        pf.space_after = Pt(1)
        pf.line_spacing = 1.0
        rr = q.add_run(re.sub(r"[`*]", "", payload))
        rr.font.size = Pt(9.5 if kind == "h2" else 8.5)
        rr.bold = kind == "h2"
        rr.font.color.rgb = ACCENT2 if kind == "h2" else GREY


def page_field(par):
    r = par.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    cached = OxmlElement("w:t")
    cached.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, cached, end):
        r._r.append(el)
    return r


def footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead = p.add_run("OptiDesk — Technical Architecture      ")
    for r in (lead, page_field(p)):
        r.font.size = Pt(8)
        r.font.color.rgb = GREY


def render(doc, blocks):
    count, last = 0, "OptiDesk"
    for kind, payload in blocks:
        if kind == "h1":
            heading(doc, 1, payload)
            subtitle(doc)
            toc(doc, blocks)
        elif kind in ("h2", "h3", "h4"):
            heading(doc, int(kind[1]), payload)
            last = re.sub(r"[`*]", "", payload)
        elif kind == "para":
            add_inline(doc.add_paragraph(), payload)
        elif kind == "bullets":
            bullets(doc, payload)
        elif kind == "numbers":
            numbers(doc, payload)
        elif kind == "table":
            table_block(doc, payload)
        elif kind == "quote":
            callout(doc, payload)
        elif kind == "code":
            code_block(doc, payload)
        elif kind == "mermaid":
            count += 1
            if not diagram(doc, count, last):
                code_block(doc, payload)
    return count


def main():
    with open(SRC, encoding="utf-8") as fh:
        blocks = parse(fh.read())
    if "--extract" in sys.argv:
        extract(blocks)
        return
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.7)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    footer(sec)
    props = doc.core_properties
    props.title = "OptiDesk — Technical Architecture Report"
    props.author = "Nishu"
    props.subject = "MERN + Clerk + Gemini AI helpdesk: architecture, API, schema, runbook"
    props.comments = "Generated from ARCHITECTURE.md"
    drawn = render(doc, blocks)
    doc.save(OUT)
    tables = sum(1 for k, _ in blocks if k == "table")
    print("wrote %s\n  %d blocks, %d tables, %d diagrams embedded"
          % (OUT, len(blocks), tables, drawn))


if __name__ == "__main__":
    main()
